from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def cell_value(cell):
    value = cell.value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


path = Path(sys.argv[1])
workbook = load_workbook(path, data_only=False, read_only=False)

result = {"file": str(path), "sheets": []}
for sheet in workbook.worksheets:
    used_rows = []
    for row in sheet.iter_rows():
        values = [cell_value(cell) for cell in row]
        if any(value not in (None, "") for value in values):
            used_rows.append(values)

    headers = used_rows[0] if used_rows else []
    records = []
    for values in used_rows[1:]:
        record = {}
        for index, header in enumerate(headers):
            key = str(header).strip() if header not in (None, "") else f"column_{index + 1}"
            record[key] = values[index] if index < len(values) else None
        records.append(record)

    duplicates = {}
    for candidate in ("Pitstop", "Name", "Outlet", "Pitstop Name", "Branch", "No_ID"):
        if candidate in headers:
            counts = Counter(str(row.get(candidate, "")).strip().upper() for row in records if row.get(candidate) not in (None, ""))
            duplicates[candidate] = sorted(
                ((name, count) for name, count in counts.items() if count > 1),
                key=lambda item: (-item[1], item[0]),
            )[:20]

    distributions = {}
    for candidate in ("Type", "Tier", "branch_status", "State", "Country", "Zone"):
        if candidate in headers:
            distributions[candidate] = Counter(
                "(blank)" if row.get(candidate) in (None, "") else str(row.get(candidate)).strip()
                for row in records
            ).most_common()

    null_counts = {
        str(header): sum(1 for row in records if row.get(str(header)) in (None, ""))
        for header in headers
    }

    result["sheets"].append(
        {
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "tables": list(sheet.tables.keys()),
            "headers": headers,
            "sample_rows": records[:12],
            "duplicates": duplicates,
            "distributions": distributions,
            "null_counts": null_counts,
            "hidden_rows": [idx for idx, dim in sheet.row_dimensions.items() if dim.hidden][:20],
            "hidden_columns": [key for key, dim in sheet.column_dimensions.items() if dim.hidden][:20],
        }
    )

print(json.dumps(result, indent=2, ensure_ascii=False))
